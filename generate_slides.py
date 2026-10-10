#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Sinh slide thuyết trình (PowerPoint) cho đồ án "Malware Analysis Platform".
~14 slide đầy đủ: giới thiệu, lý thuyết, kiến trúc, quy trình, demo, kết luận.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(BASE, "assets")
OUT = os.path.join(BASE, "SLIDE_MALWARE_ANALYSIS.pptx")

BG = RGBColor(0x0B, 0x0F, 0x1A)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
GREEN = RGBColor(0x10, 0xB9, 0x81)
RED = RGBColor(0xEF, 0x44, 0x44)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
TEXT = RGBColor(0xE5, 0xE7, 0xEB)
MUTED = RGBColor(0x9C, 0xA3, 0xAF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def textbox(slide, x, y, w, h, text, size=20, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    return tb


def title_slide(slide, text, sub=None):
    set_bg(slide)
    textbox(slide, 0.8, 2.6, 11.7, 1.2, text, 40, TEXT, True, PP_ALIGN.CENTER)
    if sub:
        textbox(slide, 0.8, 3.9, 11.7, 0.8, sub, 20, MUTED, False, PP_ALIGN.CENTER)


def content_title(slide, text):
    textbox(slide, 0.7, 0.4, 12, 0.9, text, 30, BLUE, True)
    # underline bar
    bar = slide.shapes.add_shape(1, Inches(0.7), Inches(1.25), Inches(2.5), Inches(0.04))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()


def bullet(slide, x, y, w, text, size=18, color=TEXT, level=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.level = level
    r = p.add_run()
    r.text = "• " + text if level == 0 else "   – " + text
    r.font.size = Pt(size)
    r.font.color.rgb = color


def add_pic(slide, path, x, y, width=None, height=None):
    if os.path.exists(path):
        if width:
            slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(width))
        elif height:
            slide.shapes.add_picture(path, Inches(x), Inches(y), height=Inches(height))


# ============ SLIDE 1: BÌA ============
s = add_slide()
set_bg(s)
textbox(s, 0.8, 2.2, 11.7, 1.4, "MALWARE ANALYSIS PLATFORM", 48, TEXT, True, PP_ALIGN.CENTER)
textbox(s, 0.8, 3.6, 11.7, 0.9, "Nền tảng phân tích mã độc tĩnh tự động", 24, BLUE, False, PP_ALIGN.CENTER)
textbox(s, 0.8, 4.6, 11.7, 0.8, "PE Header · YARA · Entropy · IoC", 20, MUTED, False, PP_ALIGN.CENTER)
textbox(s, 0.8, 6.2, 11.7, 0.6, "Đồ án môn Phân tích mã độc — Ngô Long — HUIT 2026", 15, MUTED, False, PP_ALIGN.CENTER)

# ============ SLIDE 2: GIỚI THIỆU ============
s = add_slide()
content_title(s, "1. Giới thiệu đề tài")
bullet(s, 0.8, 1.8, 11.5, "Bối cảnh: tấn công mạng gia tăng, mã độc là mối đe dọa hàng đầu.", 20)
bullet(s, 0.8, 2.8, 11.5, "Mục tiêu: xây dựng công cụ PHÂN TÍCH TĨNH mã độc tự động.", 20)
bullet(s, 0.8, 3.8, 11.5, "Đầu vào: tệp tin (PE, PHP, Python, PowerShell, VBScript...).", 20)
bullet(s, 0.8, 4.8, 11.5, "Đầu ra: hash, entropy, PE header, YARA, IoC, phân loại họ mã độc.", 20)
bullet(s, 0.8, 5.8, 11.5, "Triển khai: web (frontend + backend FastAPI).", 20)

# ============ SLIDE 3: PHÂN TÍCH TĨNH VS ĐỘNG ============
s = add_slide()
content_title(s, "2. Phân tích mã độc tĩnh vs động")
bullet(s, 0.8, 1.8, 11.5, "TĨNH: phân tích KHÔNG thực thi mẫu — an toàn, nhanh.", 20, GREEN)
bullet(s, 0.8, 2.6, 11.5, "   Đọc magic bytes, hash, entropy, PE header, IoC.", 18, MUTED, 1)
bullet(s, 0.8, 3.6, 11.5, "ĐỘNG: chạy mẫu trong sandbox quan sát hành vi.", 20, BLUE)
bullet(s, 0.8, 4.4, 11.5, "   Phát hiện mã tự giải mã khi chạy (mà tĩnh không thấy).", 18, MUTED, 1)
bullet(s, 0.8, 5.4, 11.5, "Đồ án này: TĨNH làm lõi + entropy để nhận biết mã packed.", 20, AMBER)

# ============ SLIDE 4: CẤU TRÚC PE ============
s = add_slide()
content_title(s, "3. Cấu trúc PE (Portable Executable)")
bullet(s, 0.7, 1.7, 5.5, "DOS Header (\"MZ\") + PE Signature", 17, TEXT)
bullet(s, 0.7, 2.3, 5.5, "File Header: machine, sections, timestamp", 17, TEXT)
bullet(s, 0.7, 2.9, 5.5, "Optional Header: Entry Point, ImageBase", 17, TEXT)
bullet(s, 0.7, 3.5, 5.5, "Section Table: .text, .data, .rsrc", 17, TEXT)
bullet(s, 0.7, 4.1, 5.5, "Import/Export Table → đoán hành vi", 17, AMBER)
add_pic(s, os.path.join(ASSETS, "pe_structure.png"), 7.0, 1.5, width=5.6)

# ============ SLIDE 5: YARA + ENTROPY + IOC ============
s = add_slide()
content_title(s, "4. YARA · Entropy · IoC")
bullet(s, 0.8, 1.7, 11.5, "YARA: quy tắc nhận diện mã độc (9 rules định nghĩa sẵn).", 20, RED)
bullet(s, 0.8, 2.6, 11.5, "Entropy: đo độ nén/mã hóa (cao > 7 = nghi packed).", 20, GREEN)
bullet(s, 0.8, 3.5, 11.5, "IoC: IP/URL/domain — dấu vết máy chủ C2.", 20, AMBER)
add_pic(s, os.path.join(ASSETS, "entropy.png"), 1.5, 4.4, width=9.5)

# ============ SLIDE 6: KIẾN TRÚC ============
s = add_slide()
content_title(s, "5. Kiến trúc hệ thống")
add_pic(s, os.path.join(ASSETS, "architecture.png"), 1.5, 1.6, width=10.3)

# ============ SLIDE 7: STACK CÔNG NGHỆ ============
s = add_slide()
content_title(s, "6. Stack công nghệ")
add_pic(s, os.path.join(ASSETS, "tech_stack.png"), 1.5, 1.8, width=10.3)

# ============ SLIDE 8: QUY TRÌNH 6 BƯỚC ============
s = add_slide()
content_title(s, "7. Quy trình phân tích (6 bước)")
add_pic(s, os.path.join(ASSETS, "pipeline.png"), 6.8, 1.3, height=5.8)

# ============ SLIDE 9: MODULE ============
s = add_slide()
content_title(s, "8. Các module phân tích")
bullet(s, 0.8, 1.8, 11.5, "malware_analyzer.py — hash, entropy, PE, IoC, phân loại họ.", 20, BLUE)
bullet(s, 0.8, 2.8, 11.5, "yara_engine.py — nạp & biên dịch YARA rules, quét mẫu.", 20, RED)
bullet(s, 0.8, 3.8, 11.5, "webshell_detector.py — regex + data-flow + polyglot.", 20, AMBER)
bullet(s, 0.8, 4.8, 11.5, "cvss_scoring.py — CVSS 3.1 + CWE + MITRE ATT&CK.", 20, PURPLE)
bullet(s, 0.8, 5.8, 11.5, "api_server.py — FastAPI nhận file upload.", 20, GREEN)

# ============ SLIDE 10: VÍ DỤ 1 ============
s = add_slide()
content_title(s, "9. Ví dụ 1: Webshell PHP")
textbox(s, 0.8, 1.7, 6, 1.5, "<?php\neval($_POST['cmd']);\n?>", 20, RED)
bullet(s, 0.8, 3.4, 11.5, "Family: Webshell (web backdoor)", 19, RED)
bullet(s, 0.8, 4.2, 11.5, "YARA: Webshell_PHP_Eval (CWE-94, T1505.003)", 19, TEXT)
bullet(s, 0.8, 5.0, 11.5, "Entropy 5.31 — mã nguồn đọc được", 19, GREEN)
bullet(s, 0.8, 5.8, 11.5, "→ Kẻ tấn công gửi lệnh qua $_POST để RCE", 19, AMBER)

# ============ SLIDE 11: VÍ DỤ 2 (OBFUSCATED) ============
s = add_slide()
content_title(s, "10. Ví dụ 2: Webshell Obfuscated (base64)")
textbox(s, 0.8, 1.7, 8, 1.5, "$a = base64_decode('c2hlbGxfZXhlYw==');\n@eval($a($_POST['x']));", 18, RED)
bullet(s, 0.8, 3.6, 11.5, "Family: PHP Obfuscated Webshell (base64/hex)", 19, RED)
bullet(s, 0.8, 4.4, 11.5, "Giải mã payload: base64_decode → \"shell_exec\"", 19, AMBER)
bullet(s, 0.8, 5.2, 11.5, "Bằng chứng: base64_decode, eval, $_POST", 19, TEXT)
bullet(s, 0.8, 6.0, 11.5, "→ Giấu lệnh độc để né quét chữ ký", 19, GREEN)

# ============ SLIDE 12: KILL CHAIN ============
s = add_slide()
content_title(s, "11. Luồng tấn công (Kill Chain)")
add_pic(s, os.path.join(ASSETS, "kill_chain.png"), 1.2, 1.4, height=5.7)

# ============ SLIDE 13: KẾT QUẢ ============
s = add_slide()
content_title(s, "12. Kết quả thực nghiệm")
bullet(s, 0.8, 1.8, 11.5, "Bộ test 45 mẫu (30 mã độc + 15 lành tính).", 20)
bullet(s, 0.8, 2.8, 11.5, "Phát hiện đúng các họ: webshell, PowerShell, VBScript, Python RAT, PE packed.", 20, GREEN)
bullet(s, 0.8, 3.8, 11.5, "Trích xuất chính xác IoC (IP/URL/domain C2).", 20)
bullet(s, 0.8, 4.8, 11.5, "Phân tích PE header + phát hiện packer UPX.", 20)
bullet(s, 0.8, 5.8, 11.5, "Sơ đồ minh họa + quy trình 6 bước phục vụ giảng dạy.", 20, AMBER)

# ============ SLIDE 14: KẾT LUẬN ============
s = add_slide()
content_title(s, "13. Kết luận & hướng phát triển")
bullet(s, 0.8, 1.8, 11.5, "Hoàn thiện nền tảng phân tích mã độc tĩnh đa định dạng.", 20, GREEN)
bullet(s, 0.8, 2.8, 11.5, "Tích hợp đủ: pefile, yara, entropy, IoC.", 20)
bullet(s, 0.8, 3.8, 11.5, "Hạn chế: chưa phân tích động, yara khó cài trên host free.", 20, AMBER)
bullet(s, 0.8, 4.8, 11.5, "Hướng đi: sandbox động, unpacker, VirusTotal API, ML.", 20, BLUE)
textbox(s, 0.8, 6.4, 11.7, 0.8, "Cảm ơn thầy cô và các bạn đã lắng nghe!", 26, TEXT, True, PP_ALIGN.CENTER)

# ============ SLIDE 15: Q&A ============
s = add_slide()
set_bg(s)
textbox(s, 0.8, 3.0, 11.7, 1.2, "Q & A", 60, GREEN, True, PP_ALIGN.CENTER)
textbox(s, 0.8, 4.3, 11.7, 0.8, "Câu hỏi & thảo luận", 22, MUTED, False, PP_ALIGN.CENTER)

prs.save(OUT)
print(f"Đã tạo slide: {OUT}")
print(f"Số slide: {len(prs.slides._sldIdLst)}")